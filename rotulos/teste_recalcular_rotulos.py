"""Teste de regressão: monta uma planilha de origem com os itens da análise de 15/10/2026
e confere se o script reproduz os totais (B2B 147 / 50 / 14 = 211).

Uso: python rotulos/teste_recalcular_rotulos.py
"""
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import recalcular_rotulos as rr  # noqa: E402

UC, C, A, P = 'Fragrância ultra concentrada', 'Fragrância concentrada', 'Aromatizador de Ambientes', 'Perfume para Ambientes'
SEM = 'Sem pedidos no período'

# (nome, concentração, status, volumes) da aba "B2B - Volumes"
IDENTIDADES = [
    ('Id Essence', C, 'Notificar', '1000ml: 4 un'),
    ('Id Ibis', A, 'Notificar', None), ('Id Ibis', C, 'Notificar', None), ('Id Ibis', 'Fragrância Ultra Concentrado', 'Notificar', None),
    ('Id Pullman', A, 'Notificar', None), ('Id Pullman', C, 'Notificar', '1000ml: 1 un'), ('Id Pullman', UC, 'Notificar', None),
    ('Id Care Plus', C, 'Notificar', None), ('Id Care Plus', UC, 'Notificar', None),
    ('ID SWEET', P, 'Notificar', None), ('ID SWEET', C, 'Notificar', '400ml: 6 un | 1000ml: 5 un'), ('ID SWEET', UC, 'Notificar', '500ml: 1 un'),
    ('Id Candy', UC, 'Notificar', '500ml: 3 un'),
    ('Id Dubai', C, 'Notificar', '400ml: 5 un'), ('Id Dubai', UC, 'Notificar', '100ml: 10 un'),
    ('ID Mercure', A, 'Notificado', None), ('ID Mercure', C, 'Notificar', '1000ml: 1 un'), ('ID Mercure', UC, 'Notificado', None),
    ('ID Novotel', A, 'Notificado', None), ('ID Novotel', C, 'Notificar', '400ml: 6 un'),
    ('ID Novotel', UC, 'Notificado', '100ml: 1 un | 500ml: 3 un | 1000ml: 2 un'),
    ('ID Grand Mercure', A, 'Notificado', None), ('ID Grand Mercure', C, 'Notificar', '400ml: 8 un | 1000ml: 2 un'),
    ('ID Grand Mercure', UC, 'Notificado', '100ml: 28 un | 1000ml: 2 un'),
    ('ID Pacco', None, 'Notificar', None), ('ID Mercedes', UC, 'Notificar', None),
    ('ID HERBAL FRESH', UC, 'Notificar', '100ml: 34 un'),
]
FRAGRANCIAS = [
    ('POMEGRANATE', P, 'Notificar', None), ('POMEGRANATE', C, 'Notificado', '400ml: 1 un'),
    ('DOCE ALECRIM', A, 'Notificar', None), ('DOCE ALECRIM', P, 'Notificado', '250ml: 8 un'),
    ('DOCE ALECRIM', C, 'Notificado', '400ml: 3 un'), ('DOCE ALECRIM', UC, 'Notificar', '100ml: 1 un | 500ml: 13 un | 1000ml: 1 un'),
    ('BAMBOO', A, 'Notificar', None), ('BAMBOO', P, 'Notificar', None),
    ('BAMBOO', C, 'Notificar', '400ml: 1 un | 1000ml: 1 un'), ('BAMBOO', UC, 'Notificar', '100ml: 4 un | 500ml: 2 un'),
    ('CHA BRANCO', A, 'Notificar', None), ('CHA BRANCO', P, 'Notificar', None), ('CHA BRANCO', C, 'Notificar', None),
    ('CHA BRANCO', UC, 'Notificado', '100ml: 2 un'),
    ('MADEIRAS NOBRES', P, 'Notificar', '250ml: 2 un'), ('MADEIRAS NOBRES', C, 'Notificar', None),
    ('MADEIRAS NOBRES', UC, 'Notificar', '100ml: 1 un | 500ml: 2 un | 1000ml: 1 un'),
    ('CITRUS FLOWERS', P, 'Notificar', None), ('CITRUS FLOWERS', C, 'Notificar', '400ml: 2 un | 1000ml: 1 un'),
    ('CITRUS FLOWERS', UC, 'Notificar', None),
    ('CAPIM LIMAO', A, 'Notificar', None), ('CAPIM LIMAO', P, 'Notificar', None), ('CAPIM LIMAO', C, 'Notificar', None),
    ('CAPIM LIMAO', UC, 'Notificar', '100ml: 5 un'),
    ('FLORESTA FRESCA', A, 'Notificar', '150ml: 3 un'), ('FLORESTA FRESCA', C, 'Notificar', '400ml: 2 un'),
    ('FLORESTA FRESCA', UC, 'Notificar', '100ml: 1 un'),
    ('NATURE', A, 'Notificar', '150ml: 1 un | 250ml: 7 un'), ('NATURE', P, 'Notificar', None),
    ('NATURE', C, 'Notificar', '400ml: 2 un'), ('NATURE', UC, 'Notificar', '100ml: 1 un'),
    ('CITRIC BATH', A, 'Notificar', '250ml: 3 un'), ('CITRIC BATH', C, 'Notificar', '1000ml: 4 un'),
    ('Flor de Madeira', None, 'Não notificar', '500ml: 4 un'),
    ('Silver', None, 'Não notificar', '400ml: 2 un | 500ml: 2 un | 1000ml: 1 un'),
    ('DOCE FIGO', P, 'Não notificar', None),
    ('LAVANDA', A, 'Não notificar', None), ('LAVANDA', P, 'Não notificar', None), ('LAVANDA', C, 'Não notificar', None),
    ('SICILIANO', UC, 'Não notificar', '500ml: 1 un'), ('LUZ', C, 'Não notificar', '400ml: 1 un'),
    ('OSK', UC, 'Não notificar', '1000ml: 1 un'), ('BERGAMOTA', UC, 'Não notificar', None),
    ('DOCE VANILLA', UC, 'Não notificar', '500ml: 2 un'),
]

INSPIRACOES = ['Sea Blue', 'Flor de Laranjeira', 'Flor de Cerejeira', 'Rosée', 'Acqua Fresh',
               'Pear & Flower', 'Wood Herbs', 'Salvegy', 'Verveine', 'Tobacco']
CLASSICOS = ['Doce Alecrim', 'Alecrim', 'Bamboo', 'Flor de Figo', 'Capim Limão', 'Nature', 'Chá Branco', 'Lavanda']
NOTIFICADOS_B2C = {
    ('"Inspirações"', '250ml'): {'Sea Blue', 'Flor de Laranjeira', 'Flor de Cerejeira', 'Rosée', 'Wood Herbs', 'Salvegy', 'Tobacco'},
    ('"Inspirações"', '500ml'): {'Sea Blue', 'Rosée', 'Acqua Fresh', 'Pear & Flower', 'Wood Herbs', 'Salvegy', 'Verveine', 'Tobacco'},
    ('"Inspirações"', '15ml'): {'Rosée', 'Acqua Fresh', 'Wood Herbs', 'Salvegy'},
    ('Classicos', '250ml'): {'Doce Alecrim', 'Capim Limão'},
    ('Classicos', '100ml'): {'Doce Alecrim'},
}
FORMATOS = [('Perfume para ambientes', '250ml'), ('Aromatizador de ambientes', '500ml'),
            ('Aromatizador de ambientes', '15ml'), ('Fragrância Concentrada', '100ml')]


def montar_origem(caminho):
    wb = openpyxl.Workbook()
    b2c = wb.active
    b2c.title = 'B2C'
    b2c.append(['Coleção', 'Concentração', 'Volume', 'Status', 'Fragrância', 'Composição', 'Notificação', 'Arte rótulo '])
    for colecao, nomes in (('"Inspirações"', INSPIRACOES), ('Classicos', CLASSICOS)):
        for conc, vol in FORMATOS:
            for n in nomes:
                st = 'Notificado' if n in NOTIFICADOS_B2C.get((colecao, vol), ()) else 'Notificar'
                b2c.append([colecao, conc, vol, st, n, None, None, None])
    b2b = wb.create_sheet('B2B')
    b2b.append(['Coleção', 'Cliente', 'Concentração', 'Volume', 'Quantidade', 'Status', 'Fragrância', 'Composição', 'Notificação', 'Arte rótulo '])
    vol = wb.create_sheet('B2B - Volumes')
    vol.append(['Cliente / Fragrância', 'Concentração', 'Status', 'Base de pedidos', 'Volumes', 'Qtd. por volume', 'Total', 'Obs.'])
    for titulo, itens in (('Identidades Olfativas', IDENTIDADES), ('Fragrâncias B2B', FRAGRANCIAS)):
        vol.append([titulo])
        for nome, conc, st, vols in itens:
            total = sum(q for _, q in rr.volumes(vols)) or None
            vol.append([nome, conc, st, None, None if vols else SEM, vols, total, None])
    wb.save(caminho)


def main():
    with tempfile.TemporaryDirectory() as d:
        origem, saida = Path(d, 'origem.xlsx'), Path(d, 'saida.xlsx')
        montar_origem(origem)
        avisos = []
        linhas = rr.montar_linhas(openpyxl.load_workbook(origem, data_only=True), avisos)
        rr.gravar(linhas, saida)
        t = rr.resumo(linhas)

        assert not avisos, avisos
        assert len(linhas) == 158, len(linhas)
        assert sum(l[0] == 'B2C' for l in linhas) == 72
        assert sum(l[0] == 'B2C' and l[7] == 'Notificado' for l in linhas) == 22
        canal = t['Canal']
        assert [canal[('B2B', st)] for st in rr.STATUS] == [147, 50, 14], canal
        esperado_vol = {  # status × volume da análise original
            'Notificar': [57, 4, 12, 32, 21, 21], 'Notificado': [31, 0, 8, 4, 3, 4], 'Não notificar': [0, 0, 0, 3, 9, 2]}
        for st, vals in esperado_vol.items():
            got = [t['Volume (ml)'][(v, st)] for v in (100, 150, 250, 400, 500, 1000)]
            assert got == vals, (st, got)
        med = {m: sum(t['Medida do rótulo'][(m, st)] for st in rr.STATUS) for m in ('13 x 5 cm', '22 x 8 cm', '28 x 10,5 cm', 'A definir')}
        assert med == {'13 x 5 cm': 88, '22 x 8 cm': 72, '28 x 10,5 cm': 27, 'A definir': 24}, med

        wb = openpyxl.load_workbook(saida)
        assert wb.sheetnames == [rr.ABA, 'Resumo']
        assert wb[rr.ABA].max_row == rr.LINHA_CABECALHO + 158 + 2
    print('OK: 158 linhas, B2B 147 / 50 / 14 = 211')


if __name__ == '__main__':
    main()
