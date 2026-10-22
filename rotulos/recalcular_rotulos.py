"""Recalcula a necessidade de rótulos a partir da planilha "Controle de Rótulos.xlsx".

Lê as abas B2C, B2B e "B2B - Volumes" da planilha de origem, gera
Fragrancias_Rotulos.xlsx (aba "Todas as Fragrâncias" com filtros + aba "Resumo"
com fórmulas) e mostra no terminal a necessidade por canal, status, volume e
medida do rótulo.

Uso:
    python rotulos/recalcular_rotulos.py "Controle de Rótulos.xlsx"
    python rotulos/recalcular_rotulos.py origem.xlsx -s saida.xlsx --manter-b2c Fragrancias_Rotulos.xlsx
"""
import argparse
import collections
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Medida do rótulo B2B por volume, conforme coluna "Volume" da aba B2B. Os demais (e todo o B2C) ficam "A definir".
MEDIDAS = {100: '13 x 5 cm', 400: '22 x 8 cm', 500: '22 x 8 cm', 1000: '28 x 10,5 cm'}
# Correspondência provável indicada na própria planilha de origem (a confirmar).
ALIAS = {'DOCE FIGO': 'FLOR DE FIGO'}
STATUS = ['Notificar', 'Notificado', 'Não notificar']
CANAIS = ['B2B', 'B2C']

COLUNAS = ['Canal do rótulo', 'Coleção / Grupo', 'Fragrância / Cliente', 'Nome na base de pedidos',
           'Concentração', 'Volume (ml)', 'Medida do rótulo', 'Status notificação', 'Nº notificação',
           'Comercialização da fragrância', 'Pedidos no período', 'Qtd. rótulos', 'Composição',
           'Arte do rótulo', 'Observação']
LARGURAS = [10, 20, 22, 30, 27, 10, 13, 14, 22, 26, 14, 10, 60, 12, 40]
ABA = 'Todas as Fragrâncias'
LINHA_CABECALHO = 4

FONTE = 'Arial'
AZUL = '1F4E78'
COR_STATUS = {'Notificar': 'FCE4D6', 'Notificado': 'E2EFDA', 'Não notificar': 'EDEDED'}
BORDA = Border(*(Side(style='thin', color='D9D9D9'),) * 4)


def chave(texto):
    """Nome sem acento, em maiúsculas e sem o prefixo "ID", para comparar entre abas."""
    s = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode().upper().strip()
    return re.sub(r'^ID\s+', '', s)


def concentracao(c):
    """Padroniza as grafias da origem ("Ultra Concentrado", "ultra concentrada" etc.)."""
    if not c:
        return 'Não informada'
    l = c.lower()
    if 'ultra' in l:
        return 'Fragrância Ultra Concentrada'
    if 'concentrad' in l:
        return 'Fragrância Concentrada'
    if 'aromatizador' in l:
        return 'Aromatizador de Ambientes'
    if 'perfume' in l:
        return 'Perfume para Ambientes'
    return c


def volumes(texto):
    """'400ml: 6 un | 1000ml: 5 un' -> [(400, 6), (1000, 5)]."""
    return [(int(v), int(q)) for v, q in re.findall(r'(\d+)\s*ml:\s*(\d+)', str(texto or ''), re.I)]


def ler_b2c(wb):
    linhas = []
    for r in wb['B2C'].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        linhas.append(dict(
            colecao=r[0].strip('"').replace('Classicos', 'Clássicos'),
            conc=concentracao(r[1]),
            vol=int(re.sub(r'\D', '', str(r[2]))),
            status=r[3],
            fragrancia=str(r[4]).strip(),
            composicao=(r[5] or '').strip(),
            notificacao=r[6],
            arte=r[7],
        ))
    return linhas


def ler_detalhes_b2b(wb):
    """Composição, nº de notificação e arte de cada cliente/fragrância × concentração (aba B2B)."""
    det, grupo = {}, None
    for r in wb['B2B'].iter_rows(min_row=2, values_only=True):
        if r[0]:
            grupo = r[0]
        if not r[1] or r[1] == 'Cliente':
            continue
        d = det.setdefault((chave(r[1]), concentracao(r[2])),
                           dict(grupo=grupo, composicao=None, notificacao=None, arte=None))
        for campo, i in (('composicao', 7), ('notificacao', 8), ('arte', 9)):
            if not d[campo] and len(r) > i and r[i]:
                d[campo] = str(r[i]).strip()
    return det


def montar_linhas(wb, avisos):
    b2c = ler_b2c(wb)
    det = ler_detalhes_b2b(wb)
    nomes_b2c = {chave(x['fragrancia']) for x in b2c}
    nomes_b2b = {chave(n) for n, _ in det} | {ALIAS.get(chave(n), chave(n)) for n, _ in det}

    linhas = []
    for x in b2c:
        com = 'B2C e B2B' if chave(x['fragrancia']) in nomes_b2b else 'Somente B2C'
        linhas.append(['B2C', x['colecao'], x['fragrancia'], x['fragrancia'], x['conc'], x['vol'],
                       'A definir', x['status'], x['notificacao'], com,
                       None, None, x['composicao'], x['arte'], None])

    grupo = 'Identidades Olfativas'
    for r in wb['B2B - Volumes'].iter_rows(min_row=2, values_only=True):
        nome, conc, status, base, descricao, qtd_vol, total, obs = (list(r) + [None] * 8)[:8]
        if nome and not status:  # linha de título de grupo
            grupo = 'Identidades Olfativas' if 'identidade' in str(nome).lower() else 'Fragrâncias B2B'
            continue
        if not nome:
            continue
        c = concentracao(conc)
        d = det.get((chave(nome), c)) or det.get((chave(nome), 'Não informada')) or {}
        if grupo == 'Identidades Olfativas':
            com = 'Somente B2B (identidade de cliente)'
        else:
            k = chave(nome)
            com = 'B2C e B2B' if ALIAS.get(k, k) in nomes_b2c else 'Somente B2B'
        pares = volumes(qtd_vol)
        if total is not None and sum(q for _, q in pares) != total:
            avisos.append(f'{nome} / {c}: soma dos volumes ({sum(q for _, q in pares)}) diferente do total ({total})')
        if not pares:
            pares, obs = [(None, 0)], obs or descricao
        for v, q in pares:
            linhas.append(['B2B', grupo, nome, base or '—', c, v, MEDIDAS.get(v, 'A definir') if v else '—',
                           status, d.get('notificacao'), com, 'Com pedido' if q else 'Sem pedido', q,
                           d.get('composicao'), d.get('arte'), obs])
    return linhas


def qtd_b2c_anteriores(caminho):
    """Quantidades B2C já preenchidas numa Fragrancias_Rotulos.xlsx anterior, por SKU."""
    ws = openpyxl.load_workbook(caminho, data_only=True)[ABA]
    qtd = {}
    for r in ws.iter_rows(min_row=LINHA_CABECALHO + 1, values_only=True):
        if r[0] == 'B2C' and r[11] is not None:
            qtd[(chave(r[2]), r[4], r[5])] = r[11]
    return qtd


def gravar(linhas, saida):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA
    ws['A1'] = 'Controle de Rótulos – todas as fragrâncias (B2C + B2B)'
    ws['A1'].font = Font(name=FONTE, bold=True, size=14, color=AZUL)
    ws['A2'] = ('Use os filtros do cabeçalho. Qtd. rótulos: B2B vem da aba "B2B - Volumes"; B2C não tem '
                'quantidade na origem — preencha as células amarelas. Medida do rótulo conforme aba B2B '
                '(100 ml 13x5 · 400/500 ml 22x8 · 1000 ml 28x10,5); demais "A definir".')
    ws['A2'].font = Font(name=FONTE, italic=True, size=9, color='595959')

    H = LINHA_CABECALHO
    for i, c in enumerate(COLUNAS, 1):
        x = ws.cell(H, i, c)
        x.font = Font(name=FONTE, bold=True, color='FFFFFF')
        x.fill = PatternFill('solid', fgColor=AZUL)
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for ri, linha in enumerate(linhas, H + 1):
        for ci, v in enumerate(linha, 1):
            c = ws.cell(ri, ci, v)
            c.font = Font(name=FONTE, size=10)
            c.border = BORDA
            c.alignment = Alignment(vertical='top')
        ws.cell(ri, 8).fill = PatternFill('solid', fgColor=COR_STATUS.get(linha[7], 'FFFFFF'))
        ws.cell(ri, 12).font = Font(name=FONTE, size=10, color='0000FF')
        if linha[0] == 'B2C':
            ws.cell(ri, 12).fill = PatternFill('solid', fgColor='FFFF00')
        arte = ws.cell(ri, 14)
        if linha[13] and str(linha[13]).startswith('http'):
            arte.hyperlink = linha[13]
            arte.value = 'Abrir arte'
            arte.font = Font(name=FONTE, size=10, color='0563C1', underline='single')
    ultima = H + len(linhas)
    tabela = Table(displayName='Fragrancias', ref=f'A{H}:{get_column_letter(len(COLUNAS))}{ultima}')
    tabela.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(tabela)
    ws.cell(ultima + 2, 11, 'Total (linhas filtradas)').font = Font(name=FONTE, bold=True)
    ws.cell(ultima + 2, 12, f'=SUBTOTAL(9,L{H + 1}:L{ultima})').font = Font(name=FONTE, bold=True)
    for i, w in enumerate(LARGURAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[H].height = 32
    ws.freeze_panes = f'D{H + 1}'
    ws.sheet_view.showGridLines = False

    gravar_resumo(wb.create_sheet('Resumo'), linhas, ultima)
    wb.save(saida)


def gravar_resumo(s, linhas, ultima):
    rng = lambda col: f"'{ABA}'!${col}${LINHA_CABECALHO + 1}:${col}${ultima}"
    s['A1'] = f'Resumo – calculado a partir da aba "{ABA}"'
    s['A1'].font = Font(name=FONTE, bold=True, size=13, color=AZUL)
    linha = [3]

    def bloco(titulo, rotulo, itens, coluna, contar=False):
        r = linha[0]
        s.cell(r, 1, titulo).font = Font(name=FONTE, bold=True)
        for i, v in enumerate([rotulo] + STATUS + ['Total'], 1):
            x = s.cell(r + 1, i, v)
            x.font = Font(name=FONTE, bold=True, color='FFFFFF')
            x.fill = PatternFill('solid', fgColor=AZUL)
            x.alignment = Alignment(horizontal='center', wrap_text=True)
        func = 'COUNTIFS' if contar else 'SUMIFS'
        soma = '' if contar else rng('L') + ','
        for i, item in enumerate(itens):
            rr = r + 2 + i
            s.cell(rr, 1, item)
            for j, st in enumerate(STATUS):
                s.cell(rr, 2 + j, f'={func}({soma}{rng(coluna)},$A{rr},{rng("H")},"{st}")')
            s.cell(rr, 5, f'=SUM(B{rr}:D{rr})').font = Font(name=FONTE, bold=True)
        rt = r + 2 + len(itens)
        s.cell(rt, 1, 'Total').font = Font(name=FONTE, bold=True)
        for col in 'BCDE':
            s[f'{col}{rt}'] = f'=SUM({col}{r + 2}:{col}{rt - 1})'
            s[f'{col}{rt}'].font = Font(name=FONTE, bold=True)
        for row in s.iter_rows(min_row=r + 1, max_row=rt, max_col=5):
            for c in row:
                c.border = BORDA
        linha[0] = rt + 2
        return rt

    t_canal = bloco('Qtd. de rótulos por canal e status', 'Canal', CANAIS, 'A')
    vols = sorted({l[5] for l in linhas if l[5]})
    bloco('Qtd. de rótulos por volume e status', 'Volume (ml)', vols, 'F')
    bloco('Qtd. de rótulos por medida e status', 'Medida do rótulo',
          sorted({l[6] for l in linhas if l[6] != '—'}, key=lambda m: (m == 'A definir', m)), 'G')

    r = linha[0]
    s.cell(r, 1, 'Cenários de seleção de volume').font = Font(name=FONTE, bold=True)
    for (rot, f), i in zip([('Somente "Notificar"', f'=B{t_canal}'),
                            ('"Notificar" + "Notificado"', f'=B{t_canal}+C{t_canal}'),
                            ('Todos', f'=E{t_canal}')], range(1, 4)):
        s.cell(r + i, 1, rot).border = BORDA
        s.cell(r + i, 2, f).border = BORDA
    linha[0] = r + 5
    bloco('Nº de linhas (SKU × volume) por canal', 'Canal', CANAIS, 'A', contar=True)

    for row in s.iter_rows(min_row=3, max_row=linha[0]):
        for c in row:
            if c.value is not None and c.font.name != FONTE:
                c.font = Font(name=FONTE)
    for i, w in enumerate([30, 14, 14, 14, 12], 1):
        s.column_dimensions[get_column_letter(i)].width = w
    s.sheet_view.showGridLines = False


def resumo(linhas):
    """Mesmos totais da aba Resumo, calculados em Python para mostrar no terminal."""
    t = collections.defaultdict(collections.Counter)
    for l in linhas:
        q = l[11] or 0
        t['Canal'][(l[0], l[7])] += q
        if l[5]:
            t['Volume (ml)'][(l[5], l[7])] += q
        if l[6] != '—':
            t['Medida do rótulo'][(l[6], l[7])] += q
    return t


def imprimir(t):
    for titulo, cont in t.items():
        itens = sorted({k for k, _ in cont}, key=lambda x: (x == 'A definir', str(x).zfill(6)))
        print(f'\n{titulo:<18}' + ''.join(f'{st:>15}' for st in STATUS) + f'{"Total":>10}')
        for it in itens:
            vals = [cont[(it, st)] for st in STATUS]
            print(f'{str(it):<18}' + ''.join(f'{v:>15}' for v in vals) + f'{sum(vals):>10}')
        vals = [sum(cont[(it, st)] for it in itens) for st in STATUS]
        print(f'{"Total":<18}' + ''.join(f'{v:>15}' for v in vals) + f'{sum(vals):>10}')
    c = t['Canal']
    b2b = {st: c[('B2B', st)] for st in STATUS}
    print(f'\nCenários B2B: só Notificar = {b2b["Notificar"]} · Notificar + Notificado = '
          f'{b2b["Notificar"] + b2b["Notificado"]} · todos = {sum(b2b.values())}')


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    p.add_argument('origem', help='planilha "Controle de Rótulos.xlsx"')
    p.add_argument('-s', '--saida', default=None,
                   help='arquivo gerado (padrão: Fragrancias_Rotulos.xlsx ao lado deste script)')
    p.add_argument('--manter-b2c', metavar='XLSX',
                   help='Fragrancias_Rotulos.xlsx anterior: reaproveita as quantidades B2C já preenchidas')
    a = p.parse_args(argv)

    avisos = []
    linhas = montar_linhas(openpyxl.load_workbook(a.origem, data_only=True), avisos)
    if a.manter_b2c:
        anteriores = qtd_b2c_anteriores(a.manter_b2c)
        for l in linhas:
            if l[0] == 'B2C':
                l[11] = anteriores.get((chave(l[2]), l[4], l[5]))
    saida = a.saida or Path(__file__).with_name('Fragrancias_Rotulos.xlsx')
    gravar(linhas, saida)

    n = collections.Counter(l[0] for l in linhas)
    print(f'{saida}: {len(linhas)} linhas ({n["B2C"]} B2C + {n["B2B"]} B2B)')
    imprimir(resumo(linhas))
    for aviso in avisos:
        print('AVISO:', aviso, file=sys.stderr)


if __name__ == '__main__':
    main()
